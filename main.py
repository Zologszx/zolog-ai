import os
import re
import json
import time
import sqlite3
import asyncio
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from aiohttp import web
from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command
from aiogram.types import (
    Message,
    CallbackQuery,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    FSInputFile,
    LabeledPrice,
    PreCheckoutQuery,
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage

from google import genai
from google.genai import types

from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image


# ============================================================
# CONFIG
# ============================================================

BOT_TOKEN = os.getenv("BOT_TOKEN", "")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
ADMIN_ID = int(os.getenv("ADMIN_ID", "0") or 0)

GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-3.6-flash"
)

GEMINI_FALLBACK_MODEL = os.getenv(
    "GEMINI_FALLBACK_MODEL",
    "gemini-3.1-flash-lite"
)

ADDZOLOG_SECRET = os.getenv("ADDZOLOG_SECRET", "")

PORT = int(os.getenv("PORT", "10000"))

MAX_BOOKS_PER_JOB = 10
MAX_FILE_SIZE = 50 * 1024 * 1024
START_GENERATIONS = 10
REFERRAL_REWARD = 1

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
BOOKS_DIR = BASE_DIR / "user_files"
OUTPUTS_DIR = BASE_DIR / "outputs"

DATA_DIR.mkdir(exist_ok=True)
BOOKS_DIR.mkdir(exist_ok=True)
OUTPUTS_DIR.mkdir(exist_ok=True)

DB_PATH = DATA_DIR / "zolog.db"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s"
)

logger = logging.getLogger("zolog-ai")


# ============================================================
# GEMINI
# ============================================================

gemini_client = None

if GEMINI_API_KEY:
    try:
        gemini_client = genai.Client(api_key=GEMINI_API_KEY)
    except Exception as e:
        logger.exception("Gemini initialization error: %s", e)


# ============================================================
# BOT / DP
# ============================================================

bot = Bot(BOT_TOKEN)
dp = Dispatcher(storage=MemoryStorage())


# ============================================================
# DATABASE
# ============================================================

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER UNIQUE,
            username TEXT,
            first_name TEXT,
            language TEXT DEFAULT 'ru',
            generations INTEGER DEFAULT 10,
            is_banned INTEGER DEFAULT 0,
            created_at TEXT,
            last_seen TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS books (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER,
            filename TEXT,
            original_name TEXT,
            title TEXT DEFAULT '',
            author TEXT DEFAULT '',
            file_path TEXT,
            file_type TEXT,
            pages INTEGER DEFAULT 0,
            extracted_text TEXT,
            created_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS materials (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER,
            material_type TEXT,
            topic TEXT,
            language TEXT,
            volume INTEGER,
            volume_type TEXT,
            options TEXT,
            status TEXT DEFAULT 'created',
            file_path TEXT,
            created_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER,
            material_id INTEGER,
            status TEXT DEFAULT 'waiting',
            progress INTEGER DEFAULT 0,
            stage TEXT,
            error TEXT,
            created_at TEXT,
            updated_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS suggestions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER,
            username TEXT,
            text TEXT,
            status TEXT DEFAULT 'new',
            admin_reply TEXT,
            created_at TEXT,
            updated_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS admins (
            telegram_id INTEGER PRIMARY KEY,
            role TEXT NOT NULL DEFAULT 'moderator',
            added_by INTEGER,
            created_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER,
            package TEXT,
            stars INTEGER,
            generations INTEGER,
            telegram_charge_id TEXT UNIQUE,
            created_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS referrals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            inviter_id INTEGER,
            invited_id INTEGER UNIQUE,
            created_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            admin_id INTEGER,
            action TEXT,
            target_id INTEGER,
            details TEXT,
            created_at TEXT
        )
    """)

    # Миграция старой базы: добавляем метаданные книг.
    existing_columns = {
        row["name"]
        for row in cur.execute("PRAGMA table_info(books)").fetchall()
    }

    if "title" not in existing_columns:
        cur.execute("ALTER TABLE books ADD COLUMN title TEXT DEFAULT ''")

    if "author" not in existing_columns:
        cur.execute("ALTER TABLE books ADD COLUMN author TEXT DEFAULT ''")

    # Первичный владелец
    if ADMIN_ID:
        cur.execute("""
            INSERT OR IGNORE INTO admins
            (telegram_id, role, added_by, created_at)
            VALUES (?, 'owner', ?, ?)
        """, (
            ADMIN_ID,
            ADMIN_ID,
            datetime.now().isoformat()
        ))

    conn.commit()
    conn.close()


# ============================================================
# USERS
# ============================================================

def ensure_user(message: Message):
    user = message.from_user

    if not user:
        return

    now = datetime.now().isoformat()

    conn = db()
    cur = conn.cursor()

    cur.execute("""
        INSERT OR IGNORE INTO users
        (telegram_id, username, first_name, generations, created_at, last_seen)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        user.id,
        user.username or "",
        user.first_name or "",
        START_GENERATIONS,
        now,
        now
    ))

    cur.execute("""
        UPDATE users
        SET username = ?,
            first_name = ?,
            last_seen = ?
        WHERE telegram_id = ?
    """, (
        user.username or "",
        user.first_name or "",
        now,
        user.id
    ))

    conn.commit()
    conn.close()


def get_user(user_id: int):
    conn = db()
    row = conn.execute(
        "SELECT * FROM users WHERE telegram_id = ?",
        (user_id,)
    ).fetchone()
    conn.close()
    return row


def get_generations(user_id: int) -> int:
    row = get_user(user_id)
    return int(row["generations"]) if row else 0


def add_generations(user_id: int, amount: int):
    conn = db()
    conn.execute("""
        UPDATE users
        SET generations = MAX(0, generations + ?)
        WHERE telegram_id = ?
    """, (amount, user_id))
    conn.commit()
    conn.close()


def set_generations(user_id: int, amount: int):
    conn = db()
    conn.execute("""
        UPDATE users
        SET generations = MAX(0, ?)
        WHERE telegram_id = ?
    """, (amount, user_id))
    conn.commit()
    conn.close()


def spend_generation(user_id: int) -> bool:
    conn = db()

    cur = conn.execute("""
        UPDATE users
        SET generations = generations - 1
        WHERE telegram_id = ?
          AND generations > 0
          AND is_banned = 0
    """, (user_id,))

    success = cur.rowcount > 0

    conn.commit()
    conn.close()

    return success


# ============================================================
# ADMINS
# ============================================================

ROLE_OWNER = "owner"
ROLE_ADMIN = "admin"
ROLE_MODERATOR = "moderator"
ROLE_ANALYST = "analyst"


def get_admin_role(user_id: int) -> Optional[str]:
    conn = db()
    row = conn.execute(
        "SELECT role FROM admins WHERE telegram_id = ?",
        (user_id,)
    ).fetchone()
    conn.close()

    if not row:
        return None

    return row["role"]


def is_admin(user_id: int) -> bool:
    return get_admin_role(user_id) is not None


def is_owner(user_id: int) -> bool:
    return get_admin_role(user_id) == ROLE_OWNER


def admin_allowed(user_id: int, minimum="moderator") -> bool:
    role = get_admin_role(user_id)

    if role == ROLE_OWNER:
        return True

    levels = {
        ROLE_ANALYST: 1,
        ROLE_MODERATOR: 2,
        ROLE_ADMIN: 3,
        ROLE_OWNER: 4
    }

    return levels.get(role, 0) >= levels.get(minimum, 2)


def log_admin(admin_id, action, target_id=None, details=""):
    conn = db()
    conn.execute("""
        INSERT INTO logs
        (admin_id, action, target_id, details, created_at)
        VALUES (?, ?, ?, ?, ?)
    """, (
        admin_id,
        action,
        target_id,
        details,
        datetime.now().isoformat()
    ))
    conn.commit()
    conn.close()


# ============================================================
# STATES
# ============================================================

class GenerationStates(StatesGroup):
    waiting_books = State()
    waiting_book_info = State()
    waiting_language = State()
    waiting_type = State()
    waiting_topic = State()
    waiting_volume = State()
    waiting_options = State()
    waiting_confirmation = State()


class SuggestionStates(StatesGroup):
    waiting_text = State()


class AskStates(StatesGroup):
    waiting_question = State()


class AdminStates(StatesGroup):
    waiting_broadcast = State()
    waiting_give_id = State()
    waiting_give_amount = State()
    waiting_add_admin_id = State()
    waiting_suggestion_reply = State()


# ============================================================
# KEYBOARDS
# ============================================================

def main_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="📝 Создать материал",
                callback_data="create_material"
            )
        ],
        [
            InlineKeyboardButton(
                text="🧠 Спросить AI",
                callback_data="ask_ai"
            ),
            InlineKeyboardButton(
                text="📖 Мои книги",
                callback_data="my_books"
            )
        ],
        [
            InlineKeyboardButton(
                text="📚 Мои материалы",
                callback_data="my_materials"
            ),
            InlineKeyboardButton(
                text="👤 Профиль",
                callback_data="profile"
            )
        ],
        [
            InlineKeyboardButton(
                text="🎁 Пригласить друга",
                callback_data="referral"
            ),
            InlineKeyboardButton(
                text="⭐ Получить генерации",
                callback_data="buy_generations"
            )
        ],
        [
            InlineKeyboardButton(
                text="🌐 Язык",
                callback_data="language"
            ),
            InlineKeyboardButton(
                text="⚙️ Настройки",
                callback_data="settings"
            )
        ],
        [
            InlineKeyboardButton(
                text="💡 Предложения по улучшению",
                callback_data="suggestions"
            )
        ],
        [
            InlineKeyboardButton(
                text="ℹ️ Помощь",
                callback_data="help"
            )
        ]
    ])


def back_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="⬅️ Главное меню",
                callback_data="main_menu"
            )
        ]
    ])


def language_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="🇺🇦 Українська",
                callback_data="gen_lang_uk"
            )
        ],
        [
            InlineKeyboardButton(
                text="🇷🇺 Русский",
                callback_data="gen_lang_ru"
            )
        ],
        [
            InlineKeyboardButton(
                text="🇬🇧 English",
                callback_data="gen_lang_en"
            )
        ]
    ])


def material_types_keyboard():
    types = [
        ("📄 Реферат", "ref"),
        ("📑 Курсовая", "course"),
        ("🎓 Дипломная", "thesis"),
        ("📋 Доклад", "report"),
        ("📚 Конспект", "summary"),
        ("📊 Презентация", "presentation"),
        ("✍️ Эссе", "essay"),
        ("📝 Свой запрос", "custom")
    ]

    rows = []

    for name, code in types:
        rows.append([
            InlineKeyboardButton(
                text=name,
                callback_data=f"gen_type_{code}"
            )
        ])

    rows.append([
        InlineKeyboardButton(
            text="⬅️ Назад",
            callback_data="create_material"
        )
    ])

    return InlineKeyboardMarkup(inline_keyboard=rows)


def confirmation_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="✅ Начать генерацию",
                callback_data="generation_start"
            )
        ],
        [
            InlineKeyboardButton(
                text="✏️ Изменить параметры",
                callback_data="generation_edit"
            )
        ],
        [
            InlineKeyboardButton(
                text="❌ Отмена",
                callback_data="main_menu"
            )
        ]
    ])


def book_upload_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="➕ Добавить книгу",
                callback_data="book_add"
            )
        ],
        [
            InlineKeyboardButton(
                text="➡️ Продолжить",
                callback_data="books_continue"
            )
        ],
        [
            InlineKeyboardButton(
                text="❌ Отмена",
                callback_data="main_menu"
            )
        ]
    ])


def presentation_options_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="🖼 Изображения: Да",
                callback_data="opt_images_yes"
            ),
            InlineKeyboardButton(
                text="🖼 Изображения: Нет",
                callback_data="opt_images_no"
            )
        ],
        [
            InlineKeyboardButton(
                text="📊 Таблицы/схемы: Да",
                callback_data="opt_tables_yes"
            ),
            InlineKeyboardButton(
                text="📊 Таблицы/схемы: Нет",
                callback_data="opt_tables_no"
            )
        ],
        [
            InlineKeyboardButton(
                text="🎤 Заметки докладчика: Да",
                callback_data="opt_notes_yes"
            ),
            InlineKeyboardButton(
                text="🎤 Заметки: Нет",
                callback_data="opt_notes_no"
            )
        ],
        [
            InlineKeyboardButton(
                text="➡️ Далее",
                callback_data="options_continue"
            )
        ]
    ])


def academic_options_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="📑 Содержание",
                callback_data="academic_toc"
            )
        ],
        [
            InlineKeyboardButton(
                text="📚 Список литературы",
                callback_data="academic_refs"
            )
        ],
        [
            InlineKeyboardButton(
                text="📊 Таблицы",
                callback_data="academic_tables"
            )
        ],
        [
            InlineKeyboardButton(
                text="📐 Схемы",
                callback_data="academic_schemes"
            )
        ],
        [
            InlineKeyboardButton(
                text="📎 Приложения",
                callback_data="academic_appendices"
            )
        ],
        [
            InlineKeyboardButton(
                text="➡️ Далее",
                callback_data="options_continue"
            )
        ]
    ])


# ============================================================
# START
# ============================================================

@dp.message(Command("start"))
async def start_handler(message: Message, state: FSMContext):
    ensure_user(message)

    await state.clear()

    args = message.text.split(maxsplit=1)

    if len(args) > 1 and args[1].startswith("ref_"):
        try:
            inviter_id = int(args[1][4:])

            if inviter_id != message.from_user.id:
                conn = db()

                existing = conn.execute("""
                    SELECT id FROM referrals
                    WHERE invited_id = ?
                """, (message.from_user.id,)).fetchone()

                if not existing:
                    conn.execute("""
                        INSERT INTO referrals
                        (inviter_id, invited_id, created_at)
                        VALUES (?, ?, ?)
                    """, (
                        inviter_id,
                        message.from_user.id,
                        datetime.now().isoformat()
                    ))

                    conn.commit()
                    conn.close()

                    add_generations(inviter_id, REFERRAL_REWARD)

                    try:
                        await bot.send_message(
                            inviter_id,
                            f"🎉 По вашей ссылке зарегистрировался новый пользователь!\n"
                            f"Вам начислено +{REFERRAL_REWARD} генерация."
                        )
                    except Exception:
                        pass

                else:
                    conn.close()

        except Exception:
            pass

    await message.answer(
        "🤖 Добро пожаловать в Zolog AI!\n\n"
        "Я помогу создавать учебные материалы "
        "на основе загруженных вами книг.\n\n"
        "📚 Главное правило:\n"
        "для академических материалов используются только "
        "загруженные вами источники.\n\n"
        f"🎁 Ваш баланс: {get_generations(message.from_user.id)} генераций.",
        reply_markup=main_menu()
    )


# ============================================================
# MAIN MENU
# ============================================================

@dp.callback_query(F.data == "main_menu")
async def main_menu_callback(callback: CallbackQuery, state: FSMContext):
    await state.clear()

    await callback.message.edit_text(
        "🤖 Zolog AI\n\nВыберите действие:",
        reply_markup=main_menu()
    )

    await callback.answer()


# ============================================================
# GENERATION — START
# ============================================================

@dp.callback_query(F.data == "create_material")
async def create_material(callback: CallbackQuery, state: FSMContext):
    await state.clear()

    await state.set_state(GenerationStates.waiting_books)
    await state.update_data(
        books=[],
        language=None,
        material_type=None,
        topic=None,
        volume=None,
        volume_type=None,
        options={}
    )

    await callback.message.edit_text(
        "📝 Создание нового материала\n\n"
        "Шаг 1 из 6 — загрузка источников.\n\n"
        "📚 Сначала отправьте мне книгу или несколько книг, "
        "на основе которых нужно создать материал.\n\n"
        f"Можно загрузить до {MAX_BOOKS_PER_JOB} книг для одной работы.\n\n"
        "Поддерживаются:\n"
        "• PDF\n"
        "• DOCX\n"
        "• TXT\n\n"
        "После загрузки книг нажмите «➡️ Продолжить».",
        reply_markup=book_upload_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "book_add")
async def book_add(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    books = data.get("books", [])

    if len(books) >= MAX_BOOKS_PER_JOB:
        await callback.answer(
            f"Можно использовать максимум {MAX_BOOKS_PER_JOB} книг.",
            show_alert=True
        )
        return

    await callback.message.answer(
        f"📖 Отправьте следующую книгу.\n\n"
        f"Загружено: {len(books)}/{MAX_BOOKS_PER_JOB}"
    )

    await callback.answer()


@dp.message(GenerationStates.waiting_books, F.document)
async def generation_book_upload(
    message: Message,
    state: FSMContext
):
    ensure_user(message)

    data = await state.get_data()
    books = data.get("books", [])

    if len(books) >= MAX_BOOKS_PER_JOB:
        await message.answer(
            f"❌ Максимум — {MAX_BOOKS_PER_JOB} книг."
        )
        return

    document = message.document

    if document.file_size and document.file_size > MAX_FILE_SIZE:
        await message.answer(
            "❌ Файл слишком большой.\n"
            "Максимальный размер — 50 МБ."
        )
        return

    original_name = document.file_name or "book"
    extension = Path(original_name).suffix.lower()

    if extension not in [".pdf", ".docx", ".txt"]:
        await message.answer(
            "❌ Этот формат не поддерживается.\n"
            "Используйте PDF, DOCX или TXT."
        )
        return

    user_dir = BOOKS_DIR / str(message.from_user.id)
    user_dir.mkdir(parents=True, exist_ok=True)

    safe_name = re.sub(
        r"[^a-zA-Zа-яА-Я0-9._-]",
        "_",
        original_name
    )

    filename = f"{int(time.time())}_{safe_name}"
    path = user_dir / filename

    try:
        file = await bot.get_file(document.file_id)
        await bot.download_file(
            file.file_path,
            destination=path
        )

        extracted_text, pages = extract_document(
            path,
            extension
        )

        if not extracted_text.strip():
            await message.answer(
                "❌ Не удалось извлечь текст из файла."
            )
            path.unlink(missing_ok=True)
            return

        conn = db()

        cur = conn.execute("""
            INSERT INTO books
            (
                telegram_id,
                filename,
                original_name,
                title,
                author,
                file_path,
                file_type,
                pages,
                extracted_text,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            message.from_user.id,
            filename,
            original_name,
            "",
            "",
            str(path),
            extension,
            pages,
            extracted_text,
            datetime.now().isoformat()
        ))

        book_id = cur.lastrowid

        conn.commit()
        conn.close()

        await state.update_data(
            pending_book_id=book_id
        )
        await state.set_state(
            GenerationStates.waiting_book_info
        )

        page_text = (
            f"📄 Страниц: {pages}"
            if extension == ".pdf"
            else
            "📄 Страницы: для этого формата не определяются автоматически"
        )

        await message.answer(
            "✅ Файл книги загружен.\n\n"
            f"📖 Файл: {original_name}\n"
            f"{page_text}\n\n"
            "Теперь укажите название и автора книги.\n\n"
            "Отправьте в формате:\n"
            "Название книги | Автор\n\n"
            "Например:\n"
            "Анатомия человека | И. И. Иванов"
        )

    except Exception as e:
        logger.exception(
            "Book upload error: %s",
            e
        )
        path.unlink(missing_ok=True)

        await message.answer(
            "❌ Не удалось обработать книгу."
        )


@dp.message(GenerationStates.waiting_book_info, F.text)
async def generation_book_info(
    message: Message,
    state: FSMContext
):
    data = await state.get_data()
    book_id = data.get("pending_book_id")

    if not book_id:
        await state.set_state(
            GenerationStates.waiting_books
        )
        await message.answer(
            "❌ Не найдено ожидающее заполнения книги. "
            "Загрузите её ещё раз."
        )
        return

    raw = message.text.strip()

    if "|" not in raw:
        await message.answer(
            "❌ Неверный формат.\n\n"
            "Напишите:\n"
            "Название книги | Автор"
        )
        return

    title, author = [
        item.strip()
        for item in raw.split("|", 1)
    ]

    if not title or not author:
        await message.answer(
            "❌ Нужно указать и название, и автора."
        )
        return

    conn = db()

    conn.execute("""
        UPDATE books
        SET title = ?,
            author = ?
        WHERE id = ?
          AND telegram_id = ?
    """, (
        title,
        author,
        book_id,
        message.from_user.id
    ))

    row = conn.execute("""
        SELECT *
        FROM books
        WHERE id = ?
          AND telegram_id = ?
    """, (
        book_id,
        message.from_user.id
    )).fetchone()

    conn.commit()
    conn.close()

    books = data.get("books", [])
    books.append(book_id)

    await state.update_data(
        books=books,
        pending_book_id=None
    )

    await state.set_state(
        GenerationStates.waiting_books
    )

    pages = int(row["pages"] or 0)

    page_text = (
        f"📄 Страниц: {pages}"
        if row["file_type"] == ".pdf"
        else
        "📄 Страницы: автоматически не определяются для этого формата"
    )

    await message.answer(
        "✅ Книга добавлена!\n\n"
        f"📖 {title}\n"
        f"✍️ {author}\n"
        f"{page_text}\n\n"
        f"📚 Загружено для этой работы: "
        f"{len(books)}/{MAX_BOOKS_PER_JOB}",
        reply_markup=book_upload_keyboard()
    )


@dp.callback_query(
    F.data == "books_continue",
    GenerationStates.waiting_books
)
async def books_continue(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()
    books = data.get("books", [])

    if not books:
        await callback.answer(
            "Сначала загрузите хотя бы одну книгу.",
            show_alert=True
        )
        return

    conn = db()
    placeholders = ",".join("?" for _ in books)
    rows = conn.execute(
        f"SELECT id, title, author FROM books WHERE id IN ({placeholders})",
        books
    ).fetchall()
    conn.close()

    if any(
        not row["title"].strip() or not row["author"].strip()
        for row in rows
    ):
        await callback.answer(
            "Сначала заполните название и автора каждой книги.",
            show_alert=True
        )
        return

    await state.set_state(
        GenerationStates.waiting_language
    )

    await callback.message.edit_text(
        "🌐 Шаг 2 из 6 — выберите язык материала:",
        reply_markup=language_keyboard()
    )

    await callback.answer()


# ============================================================
# LANGUAGE
# ============================================================

@dp.callback_query(
    F.data.startswith("gen_lang_"),
    GenerationStates.waiting_language
)
async def generation_language(
    callback: CallbackQuery,
    state: FSMContext
):
    lang_code = callback.data.replace("gen_lang_", "")

    names = {
        "ru": "🇷🇺 Русский",
        "uk": "🇺🇦 Українська",
        "en": "🇬🇧 English"
    }

    await state.update_data(language=lang_code)

    await state.set_state(
        GenerationStates.waiting_type
    )

    await callback.message.edit_text(
        "📑 Шаг 3 из 6 — выберите тип материала:",
        reply_markup=material_types_keyboard()
    )

    await callback.answer(
        f"Выбран язык: {names.get(lang_code, lang_code)}"
    )


# ============================================================
# MATERIAL TYPE
# ============================================================

@dp.callback_query(
    F.data.startswith("gen_type_"),
    GenerationStates.waiting_type
)
async def generation_type(
    callback: CallbackQuery,
    state: FSMContext
):
    material_type = callback.data.replace(
        "gen_type_",
        ""
    )

    names = {
        "ref": "Реферат",
        "course": "Курсовая",
        "thesis": "Дипломная",
        "report": "Доклад",
        "summary": "Конспект",
        "presentation": "Презентация",
        "essay": "Эссе",
        "custom": "Свой запрос"
    }

    await state.update_data(
        material_type=material_type
    )

    await state.set_state(
        GenerationStates.waiting_topic
    )

    await callback.message.edit_text(
        f"📑 Тип: {names.get(material_type)}\n\n"
        "Шаг 4 из 6.\n\n"
        "Напишите тему материала.\n\n"
        "Например:\n"
        "«Фізична терапія при ХОЗЛ»"
    )

    await callback.answer()


# ============================================================
# TOPIC
# ============================================================

@dp.message(GenerationStates.waiting_topic, F.text)
async def generation_topic(
    message: Message,
    state: FSMContext
):
    topic = message.text.strip()

    if len(topic) < 3:
        await message.answer(
            "❌ Тема слишком короткая. "
            "Напишите тему подробнее."
        )
        return

    await state.update_data(topic=topic)

    data = await state.get_data()

    material_type = data.get("material_type")

    await state.set_state(
        GenerationStates.waiting_volume
    )

    if material_type == "presentation":
        await message.answer(
            "📊 Шаг 5 из 6 — объём презентации.\n\n"
            "Напишите количество слайдов.\n\n"
            "Например: 15"
        )

    elif material_type == "summary":
        await message.answer(
            "📚 Шаг 5 из 6 — объём конспекта.\n\n"
            "Напишите примерный объём в страницах.\n\n"
            "Например: 5"
        )

    else:
        await message.answer(
            "📄 Шаг 5 из 6 — необходимый объём.\n\n"
            "Напишите количество страниц.\n\n"
            "Например: 20"
        )


# ============================================================
# VOLUME
# ============================================================

@dp.message(GenerationStates.waiting_volume, F.text)
async def generation_volume(
    message: Message,
    state: FSMContext
):
    text = message.text.strip()

    match = re.search(r"\d+", text)

    if not match:
        await message.answer(
            "❌ Укажите количество числом.\n"
            "Например: 15"
        )
        return

    volume = int(match.group())

    if volume <= 0:
        await message.answer(
            "❌ Объём должен быть больше нуля."
        )
        return

    if volume > 200:
        await message.answer(
            "❌ Слишком большой объём для одной генерации.\n"
            "Укажите до 200 страниц/слайдов."
        )
        return

    data = await state.get_data()

    material_type = data.get("material_type")

    volume_type = (
        "слайдов"
        if material_type == "presentation"
        else "страниц"
    )

    await state.update_data(
        volume=volume,
        volume_type=volume_type
    )

    await state.set_state(
        GenerationStates.waiting_options
    )

    if material_type == "presentation":
        await message.answer(
            "🎨 Дополнительные параметры презентации:",
            reply_markup=presentation_options_keyboard()
        )

    elif material_type in [
        "course",
        "thesis",
        "ref"
    ]:
        await message.answer(
            "⚙️ Дополнительные параметры работы:",
            reply_markup=academic_options_keyboard()
        )

    else:
        await show_confirmation(message, state)


# ============================================================
# OPTIONS
# ============================================================

async def show_confirmation(
    message: Message,
    state: FSMContext
):
    data = await state.get_data()

    language_names = {
        "ru": "🇷🇺 Русский",
        "uk": "🇺🇦 Українська",
        "en": "🇬🇧 English"
    }

    type_names = {
        "ref": "Реферат",
        "course": "Курсовая",
        "thesis": "Дипломная",
        "report": "Доклад",
        "summary": "Конспект",
        "presentation": "Презентация",
        "essay": "Эссе",
        "custom": "Свой запрос"
    }

    books = data.get("books", [])

    options = data.get("options", {})

    option_text = []

    for key, value in options.items():
        if value:
            option_text.append(
                f"• {key}: {value}"
            )

    if not option_text:
        option_text.append(
            "• Дополнительные параметры не выбраны"
        )

    await state.set_state(
        GenerationStates.waiting_confirmation
    )

    text = (
        "🔎 ПРОВЕРКА ПАРАМЕТРОВ\n\n"
        f"📚 Источников: {len(books)}\n"
        f"🌐 Язык: {language_names.get(data.get('language'))}\n"
        f"📑 Тип: {type_names.get(data.get('material_type'))}\n"
        f"📝 Тема: {data.get('topic')}\n"
        f"📐 Объём: {data.get('volume')} "
        f"{data.get('volume_type')}\n\n"
        "⚙️ Дополнительные параметры:\n"
        + "\n".join(option_text)
        + "\n\n"
        "⚠️ Генерация будет списана только после "
        "нажатия «Начать генерацию»."
    )

    await message.answer(
        text,
        reply_markup=confirmation_keyboard()
    )


@dp.callback_query(
    F.data.startswith("opt_"),
    GenerationStates.waiting_options
)
async def presentation_option(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()
    options = data.get("options", {})

    parts = callback.data.replace("opt_", "").split("_")

    if len(parts) >= 2:
        key = parts[0]
        value = " ".join(parts[1:])

        names = {
            "images": "Изображения",
            "tables": "Таблицы/схемы",
            "notes": "Заметки докладчика"
        }

        values = {
            "yes": "Да",
            "no": "Нет"
        }

        options[names.get(key, key)] = values.get(
            value,
            value
        )

    await state.update_data(options=options)

    await callback.answer("Сохранено")


@dp.callback_query(
    F.data.startswith("academic_"),
    GenerationStates.waiting_options
)
async def academic_option(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()
    options = data.get("options", {})

    names = {
        "academic_toc": "Содержание",
        "academic_refs": "Список литературы",
        "academic_tables": "Таблицы",
        "academic_schemes": "Схемы",
        "academic_appendices": "Приложения"
    }

    key = names.get(callback.data)

    if key:
        options[key] = "Да"

    await state.update_data(options=options)

    await callback.answer(
        f"{key}: Да"
    )


@dp.callback_query(
    F.data == "options_continue",
    GenerationStates.waiting_options
)
async def options_continue(
    callback: CallbackQuery,
    state: FSMContext
):
    await show_confirmation(
        callback.message,
        state
    )
    await callback.answer()


@dp.callback_query(
    F.data == "generation_edit",
    GenerationStates.waiting_confirmation
)
async def generation_edit(
    callback: CallbackQuery,
    state: FSMContext
):
    await state.set_state(
        GenerationStates.waiting_language
    )

    await callback.message.edit_text(
        "🌐 Выберите язык материала:",
        reply_markup=language_keyboard()
    )

    await callback.answer()


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_document(path: Path, extension: str):
    pages = []

    if extension == ".pdf":
        reader = PdfReader(str(path))

        for number, page in enumerate(
            reader.pages,
            start=1
        ):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""

            pages.append(
                f"[BOOK_PAGE:{number}]\n{text}"
            )

        return "\n\n".join(pages), len(reader.pages)

    if extension == ".docx":
        doc = Document(str(path))

        text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        )

        return (
            "[BOOK_PAGE:1]\n" + text,
            1
        )

    if extension == ".txt":
        text = path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        return (
            "[BOOK_PAGE:1]\n" + text,
            1
        )

    raise ValueError(
        "Unsupported document type"
    )


# ============================================================
# SOURCE PROCESSING
# ============================================================

def load_books(book_ids):
    conn = db()

    placeholders = ",".join(
        "?" for _ in book_ids
    )

    rows = conn.execute(
        f"""
        SELECT *
        FROM books
        WHERE id IN ({placeholders})
        ORDER BY id
        """,
        book_ids
    ).fetchall()

    conn.close()

    return rows


def split_pages(text):
    pattern = r"\[BOOK_PAGE:(\d+)\]\n"
    matches = list(re.finditer(pattern, text))

    pages = []

    for index, match in enumerate(matches):
        page_number = int(match.group(1))

        start = match.end()

        if index + 1 < len(matches):
            end = matches[index + 1].start()
        else:
            end = len(text)

        page_text = text[start:end].strip()

        pages.append(
            (page_number, page_text)
        )

    return pages


def build_source_context(books, topic, max_chars=110000):
    """
    Формирует контекст с указанием реальных страниц.
    Старается включить информацию из всех выбранных книг.
    """

    topic_words = {
        word.lower()
        for word in re.findall(
            r"[A-Za-zА-Яа-яІіЇїЄєҐґ]{4,}",
            topic
        )
    }

    blocks = []

    # Сначала собираем совпадения по теме.
    for book_index, book in enumerate(
        books,
        start=1
    ):
        pages = split_pages(
            book["extracted_text"]
        )

        selected = []

        for page_number, page_text in pages:
            lower = page_text.lower()

            score = sum(
                1
                for word in topic_words
                if word in lower
            )

            if score > 0:
                selected.append(
                    (score, page_number, page_text)
                )

        selected.sort(
            reverse=True,
            key=lambda x: x[0]
        )

        # Берём наиболее релевантные страницы.
        selected = selected[:25]

        # Если совпадений нет — всё равно берём первые страницы.
        if not selected:
            selected = [
                (0, page, text)
                for page, text in pages[:8]
            ]

        book_title = (book["title"] or book["original_name"]).strip()
        book_author = (book["author"] or "").strip()

        blocks.append(
            f"\n===== КНИГА {book_index} =====\n"
            f"Название: {book_title}\n"
            f"Автор: {book_author}\n"
            f"Файл: {book['original_name']}\n"
        )

        for _, page_number, page_text in selected:
            blocks.append(
                f"\n[КНИГА {book_index}, СТРАНИЦА {page_number}]\n"
                f"{page_text[:7000]}"
            )

    context = "\n".join(blocks)

    return context[:max_chars]


# ============================================================
# GEMINI
# ============================================================

def language_name(code):
    return {
        "ru": "русском языке",
        "uk": "украинском языке",
        "en": "English"
    }.get(code, "русском языке")


async def ai_generate(
    prompt: str,
    model: Optional[str] = None
):
    if not gemini_client:
        raise RuntimeError(
            "GEMINI_API_KEY не настроен."
        )

    model = model or GEMINI_MODEL

    def run():
        return gemini_client.models.generate_content(
            model=model,
            contents=prompt,
            config=types.GenerateContentConfig(
                temperature=0.25
            )
        )

    response = await asyncio.to_thread(run)

    if not response or not response.text:
        raise RuntimeError(
            "AI вернул пустой ответ."
        )

    return response.text


async def ask_ai_with_fallback(prompt):
    models = [
        GEMINI_MODEL,
        GEMINI_FALLBACK_MODEL
    ]

    last_error = None

    for model in models:
        for delay in [0, 4, 8]:
            try:
                if delay:
                    await asyncio.sleep(delay)

                return await ai_generate(
                    prompt,
                    model=model
                )

            except Exception as e:
                last_error = e

                logger.warning(
                    "AI error model=%s: %s",
                    model,
                    e
                )

                error_text = str(e).lower()

                temporary = any(
                    word in error_text
                    for word in [
                        "429",
                        "503",
                        "quota",
                        "rate",
                        "unavailable",
                        "timeout",
                        "deadline"
                    ]
                )

                if not temporary:
                    break

    raise RuntimeError(
        f"AI generation failed: {last_error}"
    )


# ============================================================
# ACADEMIC PROMPT
# ============================================================

def base_source_rules():
    return """
КРИТИЧЕСКИЕ ПРАВИЛА:

1. Для академического содержания используй ТОЛЬКО
   предоставленные пользователем книги.
2. Не используй интернет для фактов, определений,
   статистики, дат и научных утверждений.
3. Не добавляй сведения из собственных знаний, если их нет
   в предоставленных источниках.
4. Не придумывай авторов, книги, издательства, годы,
   страницы, статистику или факты.
5. Если информации недостаточно, прямо укажи,
   что в предоставленных источниках недостаточно данных.
6. Ссылки ставь НЕ после каждого предложения.
   Обычно достаточно одной ссылки в конце релевантного абзаца.
7. Формат ссылки: [1, с. 25].
8. Номер страницы должен существовать в предоставленном
   контексте. Никогда не придумывай страницу.
9. Не выводи служебные маркеры [BOOK_PAGE:...] и
   [КНИГА ..., СТРАНИЦА ...].
10. Не используй Markdown: никаких #, *, ``` и декоративных
    маркеров списков.
11. Не пиши список литературы внутри разделов.
12. Список литературы всегда должен находиться
    В САМОМ КОНЦЕ документа.
13. Не называй интернет-источники.
14. Не ссылайся на Wikipedia или сайты как на источники.
15. Пиши нормальным связным академическим текстом.
"""


# ============================================================
# OUTLINE
# ============================================================

async def generate_outline(
    material_type,
    topic,
    volume,
    language,
    source_context,
    options
):
    if material_type == "ref":
        structure_instruction = """
Для реферата обязательно создай:
1. Введение
2. Основные разделы по теме
3. Заключение

Каждый пункт основной части должен быть достаточно
конкретным, чтобы его можно было полностью раскрыть.
"""
    elif material_type in ("course", "thesis"):
        structure_instruction = """
Создай полноценную структуру с введением, главами,
подразделами, выводами и заключением.
"""
    elif material_type == "report":
        structure_instruction = """
Создай компактную, но полную структуру доклада.
"""
    else:
        structure_instruction = """
Создай логичный план с последовательным раскрытием темы.
"""

    prompt = f"""
Ты создаёшь академический материал.

Тип: {material_type}
Тема: {topic}
Объём: {volume}
Язык: {language_name(language)}

{base_source_rules()}

{structure_instruction}

Верни СТРОГО JSON без Markdown:

{{
  "title": "Название материала",
  "plan": [
    "Введение",
    "1. ...",
    "1.1. ...",
    "2. ...",
    "Заключение"
  ]
}}

Не пиши сам материал.

ИСТОЧНИКИ:

{source_context}
"""

    result = await ask_ai_with_fallback(prompt)

    def parse(value):
        clean = value.strip()
        clean = re.sub(
            r"^```(?:json)?\s*",
            "",
            clean,
            flags=re.IGNORECASE
        )
        clean = re.sub(
            r"\s*```$",
            "",
            clean
        )

        data = json.loads(clean)

        if (
            isinstance(data, dict)
            and isinstance(data.get("plan"), list)
        ):
            return data

        raise ValueError("Invalid outline JSON")

    try:
        return parse(result)
    except Exception:
        repaired = await ask_ai_with_fallback(
            "Исправь следующий ответ в СТРОГО JSON "
            '{"title":"...","plan":["..."]}. '
            "Не добавляй ничего кроме JSON.\n\n"
            + result
        )
        return parse(repaired)


# ============================================================
# SECTION GENERATION
# ============================================================

def safe_text(value) -> str:
    """Never let coroutine/None/object representations leak into generated files."""
    if value is None:
        return ""
    if asyncio.iscoroutine(value):
        try:
            value.close()
        except Exception:
            pass
        return ""
    return str(value)

def clean_generated_text(text: str) -> str:
    if not text:
        return ""

    text = re.sub(
        r"```(?:text|markdown)?",
        "",
        text,
        flags=re.IGNORECASE
    )
    text = text.replace("```", "")
    text = re.sub(
        r"\[BOOK_PAGE:\d+\]",
        "",
        text,
        flags=re.IGNORECASE
    )
    text = re.sub(
        r"\[КНИГА\s+\d+,\s*СТРАНИЦА\s+\d+\]",
        "",
        text,
        flags=re.IGNORECASE
    )
    text = re.sub(
        r"(?m)^\s*#{1,6}\s*",
        "",
        text
    )
    text = re.sub(
        r"(?m)^\s*[*•]\s+",
        "",
        text
    )
    text = re.sub(
        r"[ \t]{2,}",
        " ",
        text
    )

    return text.strip()


async def generate_section(
    material_type,
    topic,
    language,
    section,
    source_context,
    target_length
):
    prompt = f"""
Напиши ПОЛНОСТЬЮ РАСКРЫТЫЙ раздел академического материала.

Тип материала: {material_type}
Тема: {topic}
Язык: {language_name(language)}
Раздел: {section}
Примерный объём: {target_length} знаков.

{base_source_rules()}

Требования:
- Полностью раскрой именно этот пункт плана.
- Не ограничивайся несколькими общими предложениями.
- Используй связные академические абзацы.
- Ссылки ставь редко, обычно в конце релевантного абзаца.
- Не ставь страницу после каждого предложения.
- Не создавай список литературы.
- Не используй Markdown.
- Верни только готовый текст раздела.

ИСТОЧНИКИ:

{source_context}
"""

    result = await ask_ai_with_fallback(prompt)
    return clean_generated_text(result)


# ============================================================
# PRESENTATION
# ============================================================

async def generate_slides(
    topic,
    language,
    slide_count,
    source_context,
    options
):
    prompt = f"""
Создай структуру презентации.

Тема: {topic}
Количество слайдов: {slide_count}
Язык: {language_name(language)}

{base_source_rules()}

Верни строго JSON-массив.

Каждый элемент должен иметь:

{{
  "title": "Название слайда",
  "content": "Текст слайда",
  "speaker_notes": "Заметки докладчика",
  "image_query": "поисковый запрос для изображения",
  "scheme_query": "поисковый запрос для схемы",
  "references": ["[1, с. 20]"]
}}

Не используй Markdown.

Каждый слайд должен быть информативным,
но не перегруженным.

image_query и scheme_query — только поисковые запросы
для визуальных материалов. Они не являются источниками
академических фактов.

ИСТОЧНИКИ:

{source_context}
"""

    result = await ask_ai_with_fallback(prompt)

    try:
        clean = result.strip()

        if clean.startswith("```"):
            clean = re.sub(
                r"^```(?:json)?",
                "",
                clean
            )
            clean = re.sub(
                r"```$",
                "",
                clean
            )

        data = json.loads(clean)

        if isinstance(data, list):
            return data

    except Exception:
        pass

    # Если модель вернула невалидный JSON,
    # создаём один повторный запрос.
    repair_prompt = f"""
Преобразуй следующий ответ в корректный JSON.

Нужен только JSON-массив без Markdown.

Формат:

[
  {{
    "title": "...",
    "content": "...",
    "speaker_notes": "...",
    "image_query": "...",
    "scheme_query": "...",
    "references": []
  }}
]

Ответ:

{result}
"""

    repaired = await ask_ai_with_fallback(
        repair_prompt
    )

    repaired = repaired.strip()

    repaired = re.sub(
        r"^```(?:json)?",
        "",
        repaired
    )

    repaired = re.sub(
        r"```$",
        "",
        repaired
    )

    return json.loads(repaired)


# ============================================================
# REFERENCES
# ============================================================

def merge_page_numbers(numbers):
    numbers = sorted(set(int(x) for x in numbers))
    if not numbers:
        return ""

    ranges = []
    start = prev = numbers[0]

    for number in numbers[1:]:
        if number == prev + 1:
            prev = number
        else:
            ranges.append(
                str(start) if start == prev
                else f"{start}–{prev}"
            )
            start = prev = number

    ranges.append(
        str(start) if start == prev
        else f"{start}–{prev}"
    )

    return ", ".join(ranges)


def build_references(books, topic):
    lines = []

    topic_words = {
        word.lower()
        for word in re.findall(
            r"[A-Za-zА-Яа-яІіЇїЄєҐґ]{4,}",
            topic
        )
    }

    for index, book in enumerate(books, start=1):
        title = (book["title"] or book["original_name"]).strip()
        author = (book["author"] or "Автор не указан").strip()

        pages_text = ""

        if book["file_type"] == ".pdf":
            pages = split_pages(book["extracted_text"])
            selected = []

            for page_number, page_text in pages:
                lower = page_text.lower()
                score = sum(
                    1
                    for word in topic_words
                    if word in lower
                )

                if score > 0:
                    selected.append(page_number)

            selected = selected[:25]

            if not selected:
                selected = [
                    page_number
                    for page_number, _ in pages[:8]
                ]

            if selected:
                pages_text = (
                    "; использованные страницы: "
                    + merge_page_numbers(selected)
                )
        else:
            pages_text = (
                "; страницы автоматически не определяются "
                f"для формата {book['file_type']}"
            )

        lines.append(
            f"{index}. {author}. {title}{pages_text}."
        )

    return "\n".join(lines)


async def generate_references(
    books,
    language,
    topic=""
):
    return build_references(books, topic)


# ============================================================
# DOCX
# ============================================================

def create_docx(
    title,
    topic,
    plan,
    sections,
    references,
    options,
    output_path
):
    doc = Document()

    title_paragraph = doc.add_paragraph()
    title_run = title_paragraph.add_run(
        clean_generated_text(title)
    )
    title_run.bold = True
    title_run.font.size = Pt(18)

    topic_paragraph = doc.add_paragraph()
    topic_paragraph.add_run(
        f"Тема: {clean_generated_text(topic)}"
    ).bold = True

    doc.add_paragraph()

    if plan:
        doc.add_heading(
            "План",
            level=1
        )

        for item in plan:
            clean_item = clean_generated_text(str(item))
            if clean_item:
                doc.add_paragraph(
                    clean_item,
                    style="List Number"
                )

    for section_title, section_text in sections:
        doc.add_heading(
            clean_generated_text(section_title),
            level=1
        )

        for paragraph in re.split(
            r"\n\s*\n",
            clean_generated_text(section_text)
        ):
            paragraph = paragraph.strip()
            if paragraph:
                doc.add_paragraph(paragraph)

    if references:
        doc.add_heading(
            "Список использованных источников",
            level=1
        )

        for line in references.splitlines():
            line = clean_generated_text(line)
            if line:
                doc.add_paragraph(line)

    doc.save(output_path)


# ============================================================
# PPTX
# ============================================================

async def search_openverse_images(query, limit=8):
    if not query:
        return []

    endpoint = "https://api.openverse.org/v1/images/"

    try:
        timeout = aiohttp.ClientTimeout(total=20)

        async with aiohttp.ClientSession(
            timeout=timeout,
            headers={"User-Agent": "Zolog-AI/1.0"}
        ) as session:
            async with session.get(
                endpoint,
                params={
                    "q": query,
                    "page_size": limit
                }
            ) as response:
                if response.status != 200:
                    return []

                data = await response.json()

        candidates = []

        for item in data.get("results", []):
            if item.get("sensitive_image"):
                continue
            if item.get("watermarked"):
                continue

            url = item.get("url")
            if not url:
                continue

            width = int(item.get("width") or 0)
            height = int(item.get("height") or 0)

            if width and height and (
                width < 600 or height < 400
            ):
                continue

            candidates.append({
                "url": url,
                "landing_url": (
                    item.get("foreign_landing_url")
                    or url
                ),
                "title": item.get("title") or "",
                "creator": item.get("creator") or "",
                "license": item.get("license") or ""
            })

        return candidates

    except Exception as e:
        logger.warning(
            "Openverse search error: %s",
            e
        )
        return []


async def download_image_candidate(
    candidate,
    destination
):
    try:
        timeout = aiohttp.ClientTimeout(total=25)

        async with aiohttp.ClientSession(
            timeout=timeout,
            headers={"User-Agent": "Zolog-AI/1.0"}
        ) as session:
            async with session.get(
                candidate["url"],
                allow_redirects=True
            ) as response:
                if response.status != 200:
                    return None

                content_type = response.headers.get(
                    "Content-Type",
                    ""
                ).lower()

                if not content_type.startswith("image/"):
                    return None

                content = await response.read()

        if len(content) > 8 * 1024 * 1024:
            return None

        raw_path = destination.with_suffix(".download")
        raw_path.write_bytes(content)

        try:
            with Image.open(raw_path) as image:
                image.verify()

            with Image.open(raw_path) as image:
                width, height = image.size

                if width < 500 or height < 300:
                    raw_path.unlink(missing_ok=True)
                    return None

                if image.mode == "RGBA":
                    final_path = destination.with_suffix(".png")
                    image.save(final_path, "PNG")
                else:
                    final_path = destination.with_suffix(".jpg")
                    image = image.convert("RGB")
                    image.save(
                        final_path,
                        "JPEG",
                        quality=90
                    )

        except Exception:
            raw_path.unlink(missing_ok=True)
            return None

        raw_path.unlink(missing_ok=True)

        return {
            "path": final_path,
            "landing_url": candidate["landing_url"],
            "creator": candidate["creator"],
            "license": candidate["license"],
            "title": candidate["title"]
        }

    except Exception as e:
        logger.warning(
            "Image download error: %s",
            e
        )
        return None


async def find_and_download_image(
    query,
    image_dir,
    index
):
    candidates = await search_openverse_images(
        query,
        limit=8
    )

    for candidate_index, candidate in enumerate(
        candidates,
        start=1
    ):
        destination = image_dir / (
            f"image_{index}_{candidate_index}.jpg"
        )

        image = await download_image_candidate(
            candidate,
            destination
        )

        if image:
            return image

    return None


def add_slide_image(
    slide,
    image_info,
    left,
    top,
    width,
    height
):
    if not image_info:
        return

    try:
        slide.shapes.add_picture(
            str(image_info["path"]),
            left,
            top,
            width=width,
            height=height
        )

        source_url = image_info.get(
            "landing_url",
            ""
        )

        if source_url:
            box = slide.shapes.add_textbox(
                left,
                top + height - Inches(0.28),
                width,
                Inches(0.25)
            )

            box.text_frame.text = (
                "Источник: " + source_url[:180]
            )

            for paragraph in box.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(5)

    except Exception as e:
        logger.warning(
            "PPTX image insert error: %s",
            e
        )


async def create_pptx(
    topic,
    slides,
    output_path,
    options,
    references=""
):
    prs = Presentation()

    title_slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    title_slide.shapes.title.text = clean_generated_text(
        topic
    )

    subtitle = title_slide.placeholders[1]

    if subtitle:
        subtitle.text = "Zolog AI"

    image_dir = (
        OUTPUTS_DIR
        / f"ppt_images_{int(time.time() * 1000)}"
    )
    image_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    for slide_number, slide_data in enumerate(
        slides,
        start=1
    ):
        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        title_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(0.25),
            Inches(9.0),
            Inches(0.7)
        )

        title_box.text_frame.text = clean_generated_text(
            str(slide_data.get("title", ""))
        )

        for paragraph in title_box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(24)
                run.bold = True

        content = clean_generated_text(
            str(slide_data.get("content", ""))
        )

        references_for_slide = slide_data.get(
            "references",
            []
        )

        if references_for_slide:
            clean_refs = [
                clean_generated_text(str(item))
                for item in references_for_slide
                if str(item).strip()
            ]

            if clean_refs:
                content += "\n\n" + " ".join(
                    clean_refs
                )

        image_info = None

        if options.get("Изображения") == "Да":
            image_query = (
                slide_data.get("image_query")
                or slide_data.get("image_hint")
                or slide_data.get("title")
                or topic
            )

            image_info = await find_and_download_image(
                str(image_query),
                image_dir,
                slide_number
            )

        body_width = (
            Inches(5.0)
            if image_info
            else Inches(8.8)
        )

        body = slide.shapes.add_textbox(
            Inches(0.55),
            Inches(1.25),
            body_width,
            Inches(4.8)
        )

        body.text_frame.word_wrap = True
        body.text_frame.text = content

        for paragraph in body.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(16)

        if image_info:
            add_slide_image(
                slide,
                image_info,
                Inches(5.7),
                Inches(1.25),
                Inches(3.75),
                Inches(4.45)
            )

        if (
            options.get("Таблицы/схемы") == "Да"
            and not image_info
        ):
            scheme_query = slide_data.get(
                "scheme_query"
            )

            if scheme_query:
                scheme_info = await find_and_download_image(
                    str(scheme_query),
                    image_dir,
                    slide_number + 1000
                )

                if scheme_info:
                    add_slide_image(
                        slide,
                        scheme_info,
                        Inches(5.7),
                        Inches(1.25),
                        Inches(3.75),
                        Inches(4.45)
                    )

        notes = clean_generated_text(
            str(slide_data.get("speaker_notes", ""))
        )

        if (
            options.get("Заметки докладчика") == "Да"
            and notes
        ):
            try:
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = notes
            except Exception:
                pass

    references_slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    ref_title = references_slide.shapes.add_textbox(
        Inches(0.5),
        Inches(0.3),
        Inches(9.0),
        Inches(0.7)
    )
    ref_title.text_frame.text = (
        "Список использованных источников"
    )

    ref_body = references_slide.shapes.add_textbox(
        Inches(0.6),
        Inches(1.2),
        Inches(8.7),
        Inches(5.6)
    )

    ref_body.text_frame.text = (
        references
        or "Список источников не сформирован."
    )

    for paragraph in ref_body.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(14)

    prs.save(output_path)

    try:
        for image_file in image_dir.iterdir():
            image_file.unlink(missing_ok=True)
        image_dir.rmdir()
    except Exception:
        pass


# ============================================================
# PROGRESS
# ============================================================

async def update_job(
    job_id,
    progress,
    stage,
    status="running"
):
    conn = db()

    conn.execute("""
        UPDATE jobs
        SET progress = ?,
            stage = ?,
            status = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        progress,
        stage,
        status,
        datetime.now().isoformat(),
        job_id
    ))

    conn.commit()
    conn.close()


async def progress_message(
    message: Message,
    text: str
):
    try:
        await message.edit_text(text)
    except Exception:
        pass


# ============================================================
# GENERATION JOB
# ============================================================

async def run_generation(
    message: Message,
    state_data: dict,
    material_id: int,
    job_id: int
):
    user_id = message.from_user.id

    try:
        books = load_books(
            state_data["books"]
        )

        await update_job(
            job_id,
            5,
            "Обработка книг"
        )

        progress = await message.answer(
            "🤖 Генерация началась.\n\n"
            "▰░░░░░░░░░░ 5%\n"
            "📚 Обрабатываю книги..."
        )

        await update_job(
            job_id,
            15,
            "Книги обработаны"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰░░░░░░░░ 15%\n"
            "📚 Книги обработаны.\n"
            "🔎 Анализирую источники..."
        )

        source_context = build_source_context(
            books,
            state_data["topic"]
        )

        await update_job(
            job_id,
            25,
            "Создание плана"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰▌░░░░░░░ 25%\n"
            "🧠 Создаю структуру материала..."
        )

        material_type = state_data[
            "material_type"
        ]

        language = state_data[
            "language"
        ]

        topic = state_data[
            "topic"
        ]

        volume = state_data[
            "volume"
        ]

        options = state_data.get(
            "options",
            {}
        )

        # ====================================================
        # PRESENTATION
        # ====================================================

        if material_type == "presentation":

            slides = await generate_slides(
                topic,
                language,
                volume,
                source_context,
                options
            )

            # Если AI вернул меньше слайдов,
            # повторно не генерируем бесконечно.
            slides = slides[:volume]

            await update_job(
                job_id,
                70,
                f"Создано слайдов: {len(slides)}"
            )

            await progress_message(
                progress,
                "🤖 Генерация презентации\n\n"
                "▰▰▰▰▰▰▰░░░ 70%\n"
                f"📊 Создано слайдов: {len(slides)}\n"
                "🔍 Проверяю источники..."
            )

            references = await generate_references(
                books,
                language,
                topic
            )

            await update_job(
                job_id,
                85,
                "Создание PPTX"
            )

            await progress_message(
                progress,
                "🤖 Генерация презентации\n\n"
                "▰▰▰▰▰▰▰▰▌░ 85%\n"
                "📚 Формирую список источников..."
            )

            filename = (
                f"presentation_{user_id}_"
                f"{int(time.time())}.pptx"
            )

            output_path = OUTPUTS_DIR / filename

            await create_pptx(
                topic,
                slides,
                output_path,
                options,
                references
            )

            await update_job(
                job_id,
                97,
                "Проверка готового файла"
            )

            await progress_message(
                progress,
                "🤖 Генерация презентации\n\n"
                "▰▰▰▰▰▰▰▰▰▌ 97%\n"
                "🔍 Финальная проверка..."
            )

            await asyncio.sleep(0.5)

            await update_job(
                job_id,
                100,
                "Готово",
                status="completed"
            )

            conn = db()

            conn.execute("""
                UPDATE materials
                SET status = 'completed',
                    file_path = ?
                WHERE id = ?
            """, (
                str(output_path),
                material_id
            ))

            conn.commit()
            conn.close()

            await progress_message(
                progress,
                "✅ Презентация готова!\n\n"
                f"📊 Слайдов: {len(slides)}"
            )

            await message.answer_document(
                FSInputFile(output_path),
                caption=(
                    f"🎉 Ваша презентация готова!\n\n"
                    f"📌 {topic}\n"
                    f"📚 Использовано книг: {len(books)}"
                )
            )

            return

        # ====================================================
        # DOCUMENT MATERIAL
        # ====================================================

        outline = await generate_outline(
            material_type,
            topic,
            volume,
            language,
            source_context,
            options
        )

        await update_job(
            job_id,
            35,
            "План создан"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰▰▌░░░░░░ 35%\n"
            "📑 План создан.\n"
            "✍️ Начинаю написание разделов..."
        )

        # Извлекаем пункты плана из JSON.
        if isinstance(outline, dict):
            plan = outline.get("plan", [])
            document_title = outline.get("title") or topic
        else:
            plan = []
            document_title = topic

        raw_sections = []

        for item in plan:
            clean = clean_generated_text(str(item))
            if clean:
                raw_sections.append(clean)

        raw_sections = raw_sections[:20]

        if not raw_sections:
            raw_sections = [
                "Введение",
                "Основная часть",
                "Заключение"
            ]

        sections = []

        total = len(raw_sections)

        for index, section_title in enumerate(
            raw_sections,
            start=1
        ):
            progress_value = 35 + int(
                45 * index / total
            )

            await update_job(
                job_id,
                progress_value,
                f"Раздел {index}/{total}"
            )

            await progress_message(
                progress,
                "🤖 Генерация материала\n\n"
                f"{'▰' * max(1, progress_value // 10)}"
                f"{'░' * max(0, 10 - progress_value // 10)} "
                f"{progress_value}%\n"
                f"✍️ Раздел {index}/{total}\n"
                f"📑 {section_title}"
            )

            target_length = max(
                700,
                int(
                    volume * 1000 / total
                )
            )

            section_text = await generate_section(
                material_type,
                topic,
                language,
                section_title,
                source_context,
                target_length
            )

            sections.append(
                (
                    section_title,
                    section_text
                )
            )

        await update_job(
            job_id,
            82,
            "Формирование источников"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰▰▰▰▰▰▰▏░ 82%\n"
            "📚 Формирую список литературы..."
        )

        references = await generate_references(
            books,
            language,
            topic
        )

        await update_job(
            job_id,
            90,
            "Создание DOCX"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰▰▰▰▰▰▰▰░ 90%\n"
            "📄 Создаю DOCX..."
        )

        filename = (
            f"material_{user_id}_"
            f"{int(time.time())}.docx"
        )

        output_path = OUTPUTS_DIR / filename

        create_docx(
            title=document_title,
            topic=topic,
            plan=raw_sections,
            sections=sections,
            references=references,
            options=options,
            output_path=output_path
        )

        await update_job(
            job_id,
            97,
            "Финальная проверка"
        )

        await progress_message(
            progress,
            "🤖 Генерация материала\n\n"
            "▰▰▰▰▰▰▰▰▰▌ 97%\n"
            "🔍 Проверяю готовый файл..."
        )

        await asyncio.sleep(0.5)

        conn = db()

        conn.execute("""
            UPDATE materials
            SET status = 'completed',
                file_path = ?
            WHERE id = ?
        """, (
            str(output_path),
            material_id
        ))

        conn.commit()
        conn.close()

        await update_job(
            job_id,
            100,
            "Готово",
            status="completed"
        )

        await progress_message(
            progress,
            "✅ Материал полностью готов!"
        )

        await message.answer_document(
            FSInputFile(output_path),
            caption=(
                "🎉 Материал готов!\n\n"
                f"📌 {topic}\n"
                f"📚 Использовано книг: {len(books)}\n"
                "📄 Формат: DOCX"
            )
        )

    except Exception as e:
        logger.exception(
            "Generation job error: %s",
            e
        )

        await update_job(
            job_id,
            0,
            "Ошибка",
            status="error"
        )

        conn = db()

        conn.execute("""
            UPDATE materials
            SET status = 'error'
            WHERE id = ?
        """, (material_id,))

        conn.execute("""
            UPDATE jobs
            SET error = ?
            WHERE id = ?
        """, (
            str(e)[:2000],
            job_id
        ))

        conn.commit()
        conn.close()

        # Возвращаем генерацию при ошибке.
        add_generations(
            user_id,
            1
        )

        await message.answer(
            "❌ Во время генерации произошла ошибка.\n\n"
            "Генерация возвращена на баланс.\n\n"
            f"Техническая причина: {str(e)[:500]}\n\n"
            "Попробуйте ещё раз."
        )


# ============================================================
# START GENERATION
# ============================================================

@dp.callback_query(
    F.data == "generation_start",
    GenerationStates.waiting_confirmation
)
async def generation_start(
    callback: CallbackQuery,
    state: FSMContext
):
    user_id = callback.from_user.id

    if get_generations(user_id) <= 0:
        await callback.answer(
            "❌ У вас нет генераций.",
            show_alert=True
        )
        return

    # Списываем только здесь.
    if not spend_generation(user_id):
        await callback.answer(
            "❌ Не удалось списать генерацию.",
            show_alert=True
        )
        return

    data = await state.get_data()

    now = datetime.now().isoformat()

    conn = db()

    cur = conn.execute("""
        INSERT INTO materials
        (
            telegram_id,
            material_type,
            topic,
            language,
            volume,
            volume_type,
            options,
            status,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, 'generating', ?)
    """, (
        user_id,
        data["material_type"],
        data["topic"],
        data["language"],
        data["volume"],
        data["volume_type"],
        json.dumps(
            data.get("options", {}),
            ensure_ascii=False
        ),
        now
    ))

    material_id = cur.lastrowid

    cur = conn.execute("""
        INSERT INTO jobs
        (
            telegram_id,
            material_id,
            status,
            progress,
            stage,
            created_at,
            updated_at
        )
        VALUES (?, ?, 'running', 0, ?, ?, ?)
    """, (
        user_id,
        material_id,
        "Запуск",
        now,
        now
    ))

    job_id = cur.lastrowid

    conn.commit()
    conn.close()

    await state.clear()

    await callback.message.edit_text(
        "🚀 Запускаю генерацию...\n\n"
        "Генерация списана.\n"
        "Можете закрыть Telegram — бот продолжит работу."
    )

    await callback.answer()

    # Фоновая задача.
    asyncio.create_task(
        run_generation(
            callback.message,
            data,
            material_id,
            job_id
        )
    )


# ============================================================
# ASK AI
# ============================================================

@dp.callback_query(F.data == "ask_ai")
async def ask_ai_start(
    callback: CallbackQuery,
    state: FSMContext
):
    await state.clear()
    await state.set_state(
        AskStates.waiting_question
    )

    await callback.message.edit_text(
        "🧠 Спросить AI\n\n"
        "Напишите свой вопрос.\n\n"
        "AI сможет отвечать на основе "
        "загруженных вами книг.",
        reply_markup=back_menu()
    )

    await callback.answer()


@dp.message(AskStates.waiting_question, F.text)
async def ask_ai_question(
    message: Message,
    state: FSMContext
):
    question = message.text.strip()

    if not question:
        return

    conn = db()

    books = conn.execute("""
        SELECT *
        FROM books
        WHERE telegram_id = ?
        ORDER BY id DESC
        LIMIT 10
    """, (
        message.from_user.id,
    )).fetchall()

    conn.close()

    if not books:
        await message.answer(
            "📚 У вас пока нет загруженных книг.\n\n"
            "Сначала загрузите книгу через "
            "«📝 Создать материал»."
        )

        await state.clear()
        return

    context = build_source_context(
        books,
        question
    )

    prompt = f"""
Ответь на вопрос пользователя.

ВОПРОС:
{question}

{base_source_rules()}

Отвечай только на основании источников ниже.

ИСТОЧНИКИ:

{context}
"""

    await message.answer(
        "🧠 Анализирую загруженные источники..."
    )

    try:
        answer = await ask_ai_with_fallback(
            prompt
        )

        # Telegram limit.
        if len(answer) <= 4000:
            await message.answer(answer)
        else:
            path = OUTPUTS_DIR / (
                f"answer_{message.from_user.id}_"
                f"{int(time.time())}.txt"
            )

            path.write_text(
                answer,
                encoding="utf-8"
            )

            await message.answer_document(
                FSInputFile(path),
                caption="🧠 Ответ AI"
            )

    except Exception:
        await message.answer(
            "❌ Не удалось получить ответ от AI."
        )

    await state.clear()


# ============================================================
# BOOKS
# ============================================================

@dp.callback_query(F.data == "my_books")
async def my_books(callback: CallbackQuery):
    conn = db()

    books = conn.execute("""
        SELECT *
        FROM books
        WHERE telegram_id = ?
        ORDER BY id DESC
        LIMIT 20
    """, (
        callback.from_user.id,
    )).fetchall()

    conn.close()

    if not books:
        text = (
            "📖 Ваша библиотека пуста.\n\n"
            "Добавьте книги через "
            "«📝 Создать материал»."
        )

    else:
        lines = [
            "📖 ВАША БИБЛИОТЕКА\n"
        ]

        for book in books:
            lines.append(
                f"#{book['id']} — "
                f"{book['original_name']} "
                f"({book['pages']} стр.)"
            )

        text = "\n".join(lines)

    await callback.message.edit_text(
        text,
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# MATERIALS
# ============================================================

@dp.callback_query(F.data == "my_materials")
async def my_materials(callback: CallbackQuery):
    conn = db()

    materials = conn.execute("""
        SELECT *
        FROM materials
        WHERE telegram_id = ?
        ORDER BY id DESC
        LIMIT 20
    """, (
        callback.from_user.id,
    )).fetchall()

    conn.close()

    if not materials:
        text = "📚 У вас пока нет созданных материалов."

    else:
        lines = [
            "📚 МОИ МАТЕРИАЛЫ\n"
        ]

        for material in materials:
            status = {
                "generating": "⏳ Генерируется",
                "completed": "✅ Готов",
                "error": "❌ Ошибка"
            }.get(
                material["status"],
                material["status"]
            )

            lines.append(
                f"#{material['id']} — "
                f"{material['topic']}\n"
                f"{status}"
            )

        text = "\n\n".join(lines)

    await callback.message.edit_text(
        text,
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# PROFILE
# ============================================================

@dp.callback_query(F.data == "profile")
async def profile(callback: CallbackQuery):
    user = get_user(
        callback.from_user.id
    )

    conn = db()

    books = conn.execute(
        "SELECT COUNT(*) FROM books WHERE telegram_id = ?",
        (callback.from_user.id,)
    ).fetchone()[0]

    materials = conn.execute(
        "SELECT COUNT(*) FROM materials WHERE telegram_id = ?",
        (callback.from_user.id,)
    ).fetchone()[0]

    referrals = conn.execute(
        "SELECT COUNT(*) FROM referrals WHERE inviter_id = ?",
        (callback.from_user.id,)
    ).fetchone()[0]

    conn.close()

    text = (
        "👤 ПРОФИЛЬ\n\n"
        f"🆔 ID: {callback.from_user.id}\n"
        f"👤 Имя: {callback.from_user.first_name}\n"
        f"⭐ Генераций: {user['generations']}\n"
        f"📖 Книг: {books}\n"
        f"📚 Материалов: {materials}\n"
        f"🎁 Приглашено друзей: {referrals}"
    )

    await callback.message.edit_text(
        text,
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# REFERRAL
# ============================================================

@dp.callback_query(F.data == "referral")
async def referral(callback: CallbackQuery):
    me = await bot.get_me()

    link = (
        f"https://t.me/{me.username}"
        f"?start=ref_{callback.from_user.id}"
    )

    await callback.message.edit_text(
        "🎁 ПРИГЛАСИТЬ ДРУГА\n\n"
        f"За каждого нового пользователя вы получите "
        f"+{REFERRAL_REWARD} генерацию.\n\n"
        "Ваша ссылка:\n"
        f"{link}",
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# BUY GENERATIONS
# ============================================================

PACKAGES = {
    "basic": {
        "name": "BASIC",
        "stars": 50,
        "generations": 100
    },
    "pro": {
        "name": "PRO",
        "stars": 150,
        "generations": 500
    },
    "premium": {
        "name": "PREMIUM",
        "stars": 350,
        "generations": 1500
    }
}


@dp.callback_query(F.data == "buy_generations")
async def buy_generations(callback: CallbackQuery):
    keyboard = InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="⭐ BASIC — 50 ⭐ → 100 генераций",
                    callback_data="buy_basic"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⭐ PRO — 150 ⭐ → 500 генераций",
                    callback_data="buy_pro"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⭐ PREMIUM — 350 ⭐ → 1500 генераций",
                    callback_data="buy_premium"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⬅️ Назад",
                    callback_data="main_menu"
                )
            ]
        ]
    )

    await callback.message.edit_text(
        "⭐ ПОЛУЧИТЬ ГЕНЕРАЦИИ\n\n"
        "Выберите пакет:",
        reply_markup=keyboard
    )

    await callback.answer()


@dp.callback_query(F.data.startswith("buy_"))
async def create_invoice(callback: CallbackQuery):
    package_id = callback.data.replace(
        "buy_",
        ""
    )

    package = PACKAGES.get(package_id)

    if not package:
        await callback.answer(
            "Пакет не найден.",
            show_alert=True
        )
        return

    prices = [
        LabeledPrice(
            label=package["name"],
            amount=package["stars"]
        )
    ]

    await bot.send_invoice(
        chat_id=callback.from_user.id,
        title=f"Zolog AI — {package['name']}",
        description=(
            f"{package['generations']} генераций"
        ),
        payload=f"zolog_{package_id}",
        currency="XTR",
        prices=prices
    )

    await callback.answer()


@dp.pre_checkout_query()
async def pre_checkout(
    query: PreCheckoutQuery
):
    await query.answer(
        ok=True
    )


@dp.message(F.successful_payment)
async def successful_payment(
    message: Message
):
    payment = message.successful_payment

    charge_id = (
        payment.telegram_payment_charge_id
    )

    conn = db()

    existing = conn.execute("""
        SELECT id
        FROM payments
        WHERE telegram_charge_id = ?
    """, (
        charge_id,
    )).fetchone()

    if existing:
        conn.close()
        return

    payload = payment.invoice_payload

    package_id = payload.replace(
        "zolog_",
        ""
    )

    package = PACKAGES.get(package_id)

    if not package:
        conn.close()
        return

    conn.execute("""
        INSERT INTO payments
        (
            telegram_id,
            package,
            stars,
            generations,
            telegram_charge_id,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        message.from_user.id,
        package_id,
        payment.total_amount,
        package["generations"],
        charge_id,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()

    add_generations(
        message.from_user.id,
        package["generations"]
    )

    await message.answer(
        "🎉 Оплата успешно получена!\n\n"
        f"⭐ Пакет: {package['name']}\n"
        f"➕ Начислено: {package['generations']} генераций\n"
        f"⭐ Потрачено: {payment.total_amount} Stars\n\n"
        f"Ваш баланс: "
        f"{get_generations(message.from_user.id)}"
    )


# ============================================================
# LANGUAGE
# ============================================================

@dp.callback_query(F.data == "language")
async def language_settings(callback: CallbackQuery):
    await callback.message.edit_text(
        "🌐 Выберите язык интерфейса:",
        reply_markup=language_keyboard()
    )

    await callback.answer()


@dp.callback_query(
    F.data.in_({
        "gen_lang_ru",
        "gen_lang_uk",
        "gen_lang_en"
    })
)
async def save_language(callback: CallbackQuery):
    lang = callback.data.replace(
        "gen_lang_",
        ""
    )

    conn = db()

    conn.execute("""
        UPDATE users
        SET language = ?
        WHERE telegram_id = ?
    """, (
        lang,
        callback.from_user.id
    ))

    conn.commit()
    conn.close()

    await callback.answer(
        "Язык сохранён."
    )

    await callback.message.edit_text(
        "🌐 Язык сохранён.",
        reply_markup=back_menu()
    )


# ============================================================
# SETTINGS
# ============================================================

@dp.callback_query(F.data == "settings")
async def settings(callback: CallbackQuery):
    await callback.message.edit_text(
        "⚙️ НАСТРОЙКИ\n\n"
        "Основные настройки доступны "
        "через профиль и выбор языка.\n\n"
        "Расширенные параметры находятся "
        "в панели администратора.",
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# HELP
# ============================================================

@dp.callback_query(F.data == "help")
async def help_callback(callback: CallbackQuery):
    await callback.message.edit_text(
        "ℹ️ ПОМОЩЬ\n\n"
        "📝 Создать материал\n"
        "Загрузите до 10 книг и выберите "
        "параметры будущей работы.\n\n"
        "📖 Источники\n"
        "Материал создаётся на основе "
        "загруженных источников.\n\n"
        "🧠 Спросить AI\n"
        "Можно задать вопрос по книгам.\n\n"
        "⭐ Генерации\n"
        "Одна генерация списывается только "
        "после подтверждения запуска.\n\n"
        "💡 Предложения\n"
        "Сообщите об ошибке или предложите "
        "новую функцию.",
        reply_markup=back_menu()
    )

    await callback.answer()


# ============================================================
# SUGGESTIONS
# ============================================================

@dp.callback_query(F.data == "suggestions")
async def suggestions_start(
    callback: CallbackQuery,
    state: FSMContext
):
    await state.clear()

    await state.set_state(
        SuggestionStates.waiting_text
    )

    await callback.message.edit_text(
        "💡 ПРЕДЛОЖЕНИЯ ПО УЛУЧШЕНИЮ\n\n"
        "Напишите своё предложение, идею "
        "или сообщите об ошибке.\n\n"
        "Например:\n"
        "«Добавьте выбор оформления презентации»\n\n"
        "Ваше сообщение увидит администрация.",
        reply_markup=back_menu()
    )

    await callback.answer()


@dp.message(
    SuggestionStates.waiting_text,
    F.text
)
async def save_suggestion(
    message: Message,
    state: FSMContext
):
    text = message.text.strip()

    if len(text) < 3:
        await message.answer(
            "❌ Напишите предложение подробнее."
        )
        return

    now = datetime.now().isoformat()

    conn = db()

    cur = conn.execute("""
        INSERT INTO suggestions
        (
            telegram_id,
            username,
            text,
            status,
            created_at,
            updated_at
        )
        VALUES (?, ?, ?, 'new', ?, ?)
    """, (
        message.from_user.id,
        message.from_user.username or "",
        text,
        now,
        now
    ))

    suggestion_id = cur.lastrowid

    conn.commit()
    conn.close()

    await state.clear()

    await message.answer(
        f"✅ Предложение #{suggestion_id} отправлено!\n\n"
        "Спасибо. Администратор сможет "
        "просмотреть его в панели.",
        reply_markup=main_menu()
    )


# ============================================================
# ADMIN PANEL
# ============================================================

def admin_keyboard():
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="👥 Пользователи",
                    callback_data="admin_users"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⭐ Генерации",
                    callback_data="admin_generations"
                ),
                InlineKeyboardButton(
                    text="🛡 Администраторы",
                    callback_data="admin_admins"
                )
            ],
            [
                InlineKeyboardButton(
                    text="📚 Книги",
                    callback_data="admin_books"
                ),
                InlineKeyboardButton(
                    text="📄 Материалы",
                    callback_data="admin_materials"
                )
            ],
            [
                InlineKeyboardButton(
                    text="💳 Платежи",
                    callback_data="admin_payments"
                ),
                InlineKeyboardButton(
                    text="🎁 Рефералы",
                    callback_data="admin_referrals"
                )
            ],
            [
                InlineKeyboardButton(
                    text="💡 Предложения",
                    callback_data="admin_suggestions"
                )
            ],
            [
                InlineKeyboardButton(
                    text="📊 Статистика",
                    callback_data="admin_stats"
                ),
                InlineKeyboardButton(
                    text="📋 Логи",
                    callback_data="admin_logs"
                )
            ],
            [
                InlineKeyboardButton(
                    text="🤖 AI",
                    callback_data="admin_ai"
                )
            ],
            [
                InlineKeyboardButton(
                    text="📢 Рассылка",
                    callback_data="admin_broadcast"
                )
            ],
            [
                InlineKeyboardButton(
                    text="⬅️ Главное меню",
                    callback_data="main_menu"
                )
            ]
        ]
    )


@dp.message(Command("apanel"))
async def apanel(message: Message):
    ensure_user(message)

    if not is_admin(message.from_user.id):
        await message.answer(
            "⛔ Доступ запрещён."
        )
        return

    role = get_admin_role(
        message.from_user.id
    )

    await message.answer(
        "👑 ZOLOG AI — АДМИН-ПАНЕЛЬ\n\n"
        f"Ваша роль: {role}\n\n"
        "Выберите раздел:",
        reply_markup=admin_keyboard()
    )


# ============================================================
# AHELP
# ============================================================

@dp.message(Command("ahelp"))
async def ahelp(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer(
            "⛔ Доступ запрещён."
        )
        return

    await message.answer(
        "🛡 АДМИН-КОМАНДЫ\n\n"
        "/apanel — админ-панель\n"
        "/ahelp — список команд\n\n"
        "/users — пользователи\n"
        "/user ID — информация о пользователе\n\n"
        "/give ID N — выдать N генераций\n"
        "/take ID N — забрать N генераций\n"
        "/setgen ID N — установить баланс\n\n"
        "/admins — список администраторов\n"
        "/addadmin ID — добавить администратора\n"
        "/removeadmin ID — удалить администратора\n\n"
        "/ban ID — заблокировать\n"
        "/unban ID — разблокировать\n\n"
        "/stats — статистика\n\n"
        "👑 Команда владельца:\n"
        "/addzolog СЕКРЕТНЫЙ_КОД\n\n"
        "Секрет хранится в Render Environment."
    )


# ============================================================
# ADDZOLOG
# ============================================================

@dp.message(Command("addzolog"))
async def addzolog(message: Message):
    """
    Безопасная версия:
    /addzolog СЕКРЕТ

    Секрет НЕ хранится в коде.
    Он задаётся в Render:
    ADDZOLOG_SECRET
    """

    if not ADDZOLOG_SECRET:
        await message.answer(
            "❌ Команда владельца не настроена."
        )
        return

    parts = message.text.split(maxsplit=1)

    if len(parts) != 2:
        await message.answer(
            "❌ Использование:\n"
            "/addzolog СЕКРЕТ"
        )
        return

    supplied_secret = parts[1].strip()

    if supplied_secret != ADDZOLOG_SECRET:
        await message.answer(
            "⛔ Неверный секретный код."
        )
        return

    user_id = message.from_user.id

    conn = db()

    conn.execute("""
        INSERT INTO admins
        (
            telegram_id,
            role,
            added_by,
            created_at
        )
        VALUES (?, 'owner', ?, ?)
        ON CONFLICT(telegram_id)
        DO UPDATE SET role = 'owner'
    """, (
        user_id,
        user_id,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()

    log_admin(
        user_id,
        "owner_access_granted",
        user_id,
        "Activated via /addzolog"
    )

    await message.answer(
        "👑 Права владельца активированы.\n\n"
        "Теперь у этого аккаунта полный доступ "
        "к админ-панели Zolog AI.\n\n"
        "Открыть:\n"
        "/apanel"
    )


# ============================================================
# ADMIN — USERS
# ============================================================

@dp.message(Command("users"))
async def users_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_MODERATOR
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    conn = db()

    users = conn.execute("""
        SELECT telegram_id, username,
               first_name, generations,
               created_at
        FROM users
        ORDER BY id DESC
        LIMIT 30
    """).fetchall()

    conn.close()

    if not users:
        await message.answer(
            "Пользователей нет."
        )
        return

    lines = ["👥 ПОСЛЕДНИЕ ПОЛЬЗОВАТЕЛИ\n"]

    for user in users:
        lines.append(
            f"🆔 {user['telegram_id']}\n"
            f"👤 @{user['username'] or 'нет'}\n"
            f"⭐ {user['generations']}"
        )

    await message.answer(
        "\n\n".join(lines)
    )


@dp.message(Command("user"))
async def user_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_MODERATOR
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) < 2:
        await message.answer(
            "Использование:\n/user ID"
        )
        return

    try:
        user_id = int(parts[1])
    except ValueError:
        await message.answer(
            "❌ Неверный ID."
        )
        return

    user = get_user(user_id)

    if not user:
        await message.answer(
            "❌ Пользователь не найден."
        )
        return

    conn = db()

    books = conn.execute(
        "SELECT COUNT(*) FROM books WHERE telegram_id = ?",
        (user_id,)
    ).fetchone()[0]

    materials = conn.execute(
        "SELECT COUNT(*) FROM materials WHERE telegram_id = ?",
        (user_id,)
    ).fetchone()[0]

    conn.close()

    await message.answer(
        "👤 ПОЛЬЗОВАТЕЛЬ\n\n"
        f"ID: {user_id}\n"
        f"Username: @{user['username'] or 'нет'}\n"
        f"Имя: {user['first_name']}\n"
        f"Генерации: {user['generations']}\n"
        f"Книги: {books}\n"
        f"Материалы: {materials}\n"
        f"Регистрация: {user['created_at']}"
    )


# ============================================================
# GENERATION MANAGEMENT
# ============================================================

@dp.message(Command("give"))
async def give_generations(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ADMIN
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Использование:\n"
            "/give ID количество"
        )
        return

    try:
        target = int(parts[1])
        amount = int(parts[2])
    except ValueError:
        await message.answer(
            "❌ Неверные значения."
        )
        return

    add_generations(
        target,
        amount
    )

    log_admin(
        message.from_user.id,
        "give_generations",
        target,
        f"+{amount}"
    )

    await message.answer(
        f"✅ Пользователю {target} выдано "
        f"+{amount} генераций.\n\n"
        f"Новый баланс: {get_generations(target)}"
    )


@dp.message(Command("take"))
async def take_generations(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ADMIN
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Использование:\n"
            "/take ID количество"
        )
        return

    try:
        target = int(parts[1])
        amount = int(parts[2])
    except ValueError:
        await message.answer(
            "❌ Неверные значения."
        )
        return

    add_generations(
        target,
        -abs(amount)
    )

    log_admin(
        message.from_user.id,
        "take_generations",
        target,
        f"-{amount}"
    )

    await message.answer(
        f"✅ У пользователя {target} забрано "
        f"{amount} генераций.\n\n"
        f"Новый баланс: {get_generations(target)}"
    )


@dp.message(Command("setgen"))
async def setgen_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ADMIN
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) != 3:
        await message.answer(
            "Использование:\n"
            "/setgen ID количество"
        )
        return

    try:
        target = int(parts[1])
        amount = int(parts[2])
    except ValueError:
        await message.answer(
            "❌ Неверные значения."
        )
        return

    set_generations(
        target,
        amount
    )

    log_admin(
        message.from_user.id,
        "set_generations",
        target,
        str(amount)
    )

    await message.answer(
        f"✅ Баланс пользователя {target} "
        f"установлен на {amount}."
    )


# ============================================================
# ADMIN MANAGEMENT
# ============================================================

@dp.message(Command("admins"))
async def admins_command(message: Message):
    if not is_admin(
        message.from_user.id
    ):
        await message.answer(
            "⛔ Доступ запрещён."
        )
        return

    conn = db()

    admins = conn.execute("""
        SELECT *
        FROM admins
        ORDER BY created_at
    """).fetchall()

    conn.close()

    lines = [
        "🛡 АДМИНИСТРАТОРЫ\n"
    ]

    for admin in admins:
        lines.append(
            f"🆔 {admin['telegram_id']}\n"
            f"Роль: {admin['role']}"
        )

    await message.answer(
        "\n\n".join(lines)
    )


@dp.message(Command("addadmin"))
async def addadmin_command(message: Message):
    if not is_owner(
        message.from_user.id
    ):
        await message.answer(
            "⛔ Только владелец может "
            "добавлять администраторов."
        )
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "Использование:\n/addadmin ID"
        )
        return

    try:
        target = int(parts[1])
    except ValueError:
        await message.answer(
            "❌ Неверный ID."
        )
        return

    conn = db()

    conn.execute("""
        INSERT INTO admins
        (
            telegram_id,
            role,
            added_by,
            created_at
        )
        VALUES (?, 'admin', ?, ?)
        ON CONFLICT(telegram_id)
        DO UPDATE SET role = 'admin'
    """, (
        target,
        message.from_user.id,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()

    log_admin(
        message.from_user.id,
        "add_admin",
        target
    )

    await message.answer(
        f"🛡 Пользователь {target} "
        f"назначен администратором."
    )


@dp.message(Command("removeadmin"))
async def removeadmin_command(message: Message):
    if not is_owner(
        message.from_user.id
    ):
        await message.answer(
            "⛔ Только владелец."
        )
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "Использование:\n/removeadmin ID"
        )
        return

    try:
        target = int(parts[1])
    except ValueError:
        await message.answer(
            "❌ Неверный ID."
        )
        return

    if target == message.from_user.id:
        await message.answer(
            "❌ Нельзя удалить самого себя."
        )
        return

    conn = db()

    conn.execute(
        "DELETE FROM admins WHERE telegram_id = ?",
        (target,)
    )

    conn.commit()
    conn.close()

    log_admin(
        message.from_user.id,
        "remove_admin",
        target
    )

    await message.answer(
        f"✅ Администратор {target} удалён."
    )


# ============================================================
# BAN
# ============================================================

@dp.message(Command("ban"))
async def ban_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ADMIN
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "/ban ID"
        )
        return

    target = int(parts[1])

    conn = db()

    conn.execute("""
        UPDATE users
        SET is_banned = 1
        WHERE telegram_id = ?
    """, (target,))

    conn.commit()
    conn.close()

    log_admin(
        message.from_user.id,
        "ban",
        target
    )

    await message.answer(
        f"🚫 Пользователь {target} заблокирован."
    )


@dp.message(Command("unban"))
async def unban_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ADMIN
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "/unban ID"
        )
        return

    target = int(parts[1])

    conn = db()

    conn.execute("""
        UPDATE users
        SET is_banned = 0
        WHERE telegram_id = ?
    """, (target,))

    conn.commit()
    conn.close()

    log_admin(
        message.from_user.id,
        "unban",
        target
    )

    await message.answer(
        f"✅ Пользователь {target} разблокирован."
    )


# ============================================================
# ADMIN CALLBACKS
# ============================================================

@dp.callback_query(F.data == "admin_users")
async def admin_users(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_MODERATOR
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    count = conn.execute(
        "SELECT COUNT(*) FROM users"
    ).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        f"👥 Пользователи\n\n"
        f"Всего: {count}\n\n"
        "Подробный поиск:\n"
        "/users\n"
        "/user ID",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_generations")
async def admin_generations(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ADMIN
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    await callback.message.edit_text(
        "⭐ УПРАВЛЕНИЕ ГЕНЕРАЦИЯМИ\n\n"
        "/give ID N — выдать\n"
        "/take ID N — забрать\n"
        "/setgen ID N — установить\n\n"
        "Можно выдавать генерации себе "
        "или любому пользователю.",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_admins")
async def admin_admins(callback: CallbackQuery):
    if not is_admin(
        callback.from_user.id
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    await callback.message.edit_text(
        "🛡 АДМИНИСТРАТОРЫ\n\n"
        "/admins — список\n"
        "/addadmin ID — добавить\n"
        "/removeadmin ID — удалить\n\n"
        "👑 Владелец имеет полный доступ.",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_books")
async def admin_books(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_MODERATOR
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    total = conn.execute(
        "SELECT COUNT(*) FROM books"
    ).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        f"📚 КНИГИ\n\n"
        f"Всего загружено: {total}",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_materials")
async def admin_materials(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_MODERATOR
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    total = conn.execute(
        "SELECT COUNT(*) FROM materials"
    ).fetchone()[0]

    generating = conn.execute("""
        SELECT COUNT(*)
        FROM materials
        WHERE status = 'generating'
    """).fetchone()[0]

    completed = conn.execute("""
        SELECT COUNT(*)
        FROM materials
        WHERE status = 'completed'
    """).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        "📄 МАТЕРИАЛЫ\n\n"
        f"Всего: {total}\n"
        f"⏳ Генерируется: {generating}\n"
        f"✅ Готово: {completed}",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_payments")
async def admin_payments(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ADMIN
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    total = conn.execute("""
        SELECT COALESCE(SUM(stars), 0)
        FROM payments
    """).fetchone()[0]

    payments = conn.execute(
        "SELECT COUNT(*) FROM payments"
    ).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        "💳 ПЛАТЕЖИ\n\n"
        f"Платежей: {payments}\n"
        f"Stars получено: {total}",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.callback_query(F.data == "admin_referrals")
async def admin_referrals(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ADMIN
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    count = conn.execute(
        "SELECT COUNT(*) FROM referrals"
    ).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        f"🎁 РЕФЕРАЛЫ\n\n"
        f"Всего приглашений: {count}",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


# ============================================================
# ADMIN SUGGESTIONS
# ============================================================

@dp.callback_query(F.data == "admin_suggestions")
async def admin_suggestions(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_MODERATOR
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    total = conn.execute(
        "SELECT COUNT(*) FROM suggestions"
    ).fetchone()[0]

    new = conn.execute("""
        SELECT COUNT(*)
        FROM suggestions
        WHERE status = 'new'
    """).fetchone()[0]

    viewed = conn.execute("""
        SELECT COUNT(*)
        FROM suggestions
        WHERE status = 'viewed'
    """).fetchone()[0]

    working = conn.execute("""
        SELECT COUNT(*)
        FROM suggestions
        WHERE status = 'working'
    """).fetchone()[0]

    implemented = conn.execute("""
        SELECT COUNT(*)
        FROM suggestions
        WHERE status = 'implemented'
    """).fetchone()[0]

    rejected = conn.execute("""
        SELECT COUNT(*)
        FROM suggestions
        WHERE status = 'rejected'
    """).fetchone()[0]

    conn.close()

    await callback.message.edit_text(
        "💡 ПРЕДЛОЖЕНИЯ\n\n"
        f"Всего: {total}\n"
        f"🆕 Новые: {new}\n"
        f"👀 Просмотрены: {viewed}\n"
        f"🔨 В работе: {working}\n"
        f"✅ Реализованы: {implemented}\n"
        f"❌ Отклонены: {rejected}\n\n"
        "Команда просмотра:\n"
        "/suggestions",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


@dp.message(Command("suggestions"))
async def suggestions_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_MODERATOR
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    conn = db()

    suggestions = conn.execute("""
        SELECT *
        FROM suggestions
        ORDER BY id DESC
        LIMIT 30
    """).fetchall()

    conn.close()

    if not suggestions:
        await message.answer(
            "💡 Предложений пока нет."
        )
        return

    lines = [
        "💡 ВСЕ ПРЕДЛОЖЕНИЯ\n"
    ]

    status_names = {
        "new": "🆕 Новое",
        "viewed": "👀 Просмотрено",
        "working": "🔨 В работе",
        "implemented": "✅ Реализовано",
        "rejected": "❌ Отклонено"
    }

    for item in suggestions:
        lines.append(
            f"#{item['id']} — "
            f"{status_names.get(item['status'], item['status'])}\n"
            f"ID: {item['telegram_id']}\n"
            f"@{item['username'] or 'нет'}\n"
            f"{item['text'][:300]}"
        )

    await message.answer(
        "\n\n".join(lines)
    )


# ============================================================
# ADMIN STATS
# ============================================================

@dp.message(Command("stats"))
async def stats_command(message: Message):
    if not admin_allowed(
        message.from_user.id,
        ROLE_ANALYST
    ):
        await message.answer(
            "⛔ Недостаточно прав."
        )
        return

    conn = db()

    users = conn.execute(
        "SELECT COUNT(*) FROM users"
    ).fetchone()[0]

    books = conn.execute(
        "SELECT COUNT(*) FROM books"
    ).fetchone()[0]

    materials = conn.execute(
        "SELECT COUNT(*) FROM materials"
    ).fetchone()[0]

    suggestions = conn.execute(
        "SELECT COUNT(*) FROM suggestions"
    ).fetchone()[0]

    generations = conn.execute(
        "SELECT COALESCE(SUM(generations), 0) FROM users"
    ).fetchone()[0]

    stars = conn.execute(
        "SELECT COALESCE(SUM(stars), 0) FROM payments"
    ).fetchone()[0]

    referrals = conn.execute(
        "SELECT COUNT(*) FROM referrals"
    ).fetchone()[0]

    conn.close()

    await message.answer(
        "📊 СТАТИСТИКА ZOLOG AI\n\n"
        f"👥 Пользователи: {users}\n"
        f"📖 Книги: {books}\n"
        f"📄 Материалы: {materials}\n"
        f"⭐ Генераций на балансах: {generations}\n"
        f"💳 Stars: {stars}\n"
        f"🎁 Рефералы: {referrals}\n"
        f"💡 Предложения: {suggestions}"
    )


@dp.callback_query(F.data == "admin_stats")
async def admin_stats(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ANALYST
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    await callback.message.edit_text(
        "📊 СТАТИСТИКА\n\n"
        "Используйте:\n"
        "/stats",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


# ============================================================
# LOGS
# ============================================================

@dp.callback_query(F.data == "admin_logs")
async def admin_logs(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ADMIN
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    conn = db()

    logs = conn.execute("""
        SELECT *
        FROM logs
        ORDER BY id DESC
        LIMIT 20
    """).fetchall()

    conn.close()

    if not logs:
        text = "📋 Логов пока нет."

    else:
        lines = [
            "📋 ПОСЛЕДНИЕ ДЕЙСТВИЯ\n"
        ]

        for log in logs:
            lines.append(
                f"{log['created_at']}\n"
                f"Admin: {log['admin_id']}\n"
                f"{log['action']}\n"
                f"Target: {log['target_id'] or '-'}\n"
                f"{log['details'] or ''}"
            )

        text = "\n\n".join(lines)

    await callback.message.edit_text(
        text,
        reply_markup=admin_keyboard()
    )

    await callback.answer()


# ============================================================
# ADMIN AI
# ============================================================

@dp.callback_query(F.data == "admin_ai")
async def admin_ai(callback: CallbackQuery):
    if not admin_allowed(
        callback.from_user.id,
        ROLE_ADMIN
    ):
        await callback.answer(
            "Недостаточно прав.",
            show_alert=True
        )
        return

    status = (
        "🟢 API key настроен"
        if GEMINI_API_KEY
        else "🔴 API key отсутствует"
    )

    await callback.message.edit_text(
        "🤖 AI СИСТЕМА\n\n"
        f"{status}\n\n"
        f"Основная модель:\n"
        f"{GEMINI_MODEL}\n\n"
        f"Fallback:\n"
        f"{GEMINI_FALLBACK_MODEL}\n\n"
        "При временных ошибках используется "
        "повторная попытка и fallback-модель.",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


# ============================================================
# BROADCAST
# ============================================================

@dp.callback_query(F.data == "admin_broadcast")
async def admin_broadcast(
    callback: CallbackQuery,
    state: FSMContext
):
    if not is_owner(
        callback.from_user.id
    ):
        await callback.answer(
            "Только владелец.",
            show_alert=True
        )
        return

    await state.set_state(
        AdminStates.waiting_broadcast
    )

    await callback.message.edit_text(
        "📢 РАССЫЛКА\n\n"
        "Отправьте сообщение, которое нужно "
        "разослать пользователям.\n\n"
        "После отправки оно будет использовано "
        "как текст рассылки.",
        reply_markup=back_menu()
    )

    await callback.answer()


@dp.message(
    AdminStates.waiting_broadcast,
    F.text
)
async def process_broadcast(
    message: Message,
    state: FSMContext
):
    if not is_owner(
        message.from_user.id
    ):
        return

    text = message.text

    conn = db()

    users = conn.execute("""
        SELECT telegram_id
        FROM users
        WHERE is_banned = 0
    """).fetchall()

    conn.close()

    await message.answer(
        f"📢 Начинаю рассылку.\n"
        f"Получателей: {len(users)}"
    )

    success = 0

    for user in users:
        try:
            await bot.send_message(
                user["telegram_id"],
                text
            )

            success += 1

            await asyncio.sleep(
                0.05
            )

        except Exception:
            pass

    await message.answer(
        f"✅ Рассылка завершена.\n\n"
        f"Успешно: {success}\n"
        f"Всего: {len(users)}"
    )

    log_admin(
        message.from_user.id,
        "broadcast",
        details=f"success={success}"
    )

    await state.clear()


# ============================================================
# ADMIN PANEL CALLBACK
# ============================================================

@dp.callback_query(F.data == "admin_back")
async def admin_back(callback: CallbackQuery):
    if not is_admin(
        callback.from_user.id
    ):
        await callback.answer(
            "Нет доступа.",
            show_alert=True
        )
        return

    await callback.message.edit_text(
        "👑 ZOLOG AI — АДМИН-ПАНЕЛЬ",
        reply_markup=admin_keyboard()
    )

    await callback.answer()


# ============================================================
# GENERATION PROGRESS COMMAND
# ============================================================

@dp.message(Command("job"))
async def job_command(message: Message):
    parts = message.text.split()

    if len(parts) != 2:
        await message.answer(
            "Использование:\n/job ID"
        )
        return

    try:
        job_id = int(parts[1])
    except ValueError:
        await message.answer(
            "❌ Неверный ID."
        )
        return

    conn = db()

    job = conn.execute("""
        SELECT *
        FROM jobs
        WHERE id = ?
          AND telegram_id = ?
    """, (
        job_id,
        message.from_user.id
    )).fetchone()

    conn.close()

    if not job:
        await message.answer(
            "❌ Задача не найдена."
        )
        return

    await message.answer(
        "📊 СТАТУС ГЕНЕРАЦИИ\n\n"
        f"ID: {job['id']}\n"
        f"Статус: {job['status']}\n"
        f"Прогресс: {job['progress']}%\n"
        f"Этап: {job['stage']}"
    )


# ============================================================
# BAN CHECK
# ============================================================

@dp.message()
async def general_message_handler(
    message: Message
):
    ensure_user(message)

    user = get_user(
        message.from_user.id
    )

    if user and user["is_banned"]:
        await message.answer(
            "🚫 Ваш аккаунт заблокирован."
        )
        return

    if message.text and message.text.startswith("/"):
        return

    if message.document:
        await message.answer(
            "📚 Чтобы добавить книгу для генерации, "
            "нажмите «📝 Создать материал»."
        )

    elif message.text:
        await message.answer(
            "Выберите действие в меню:",
            reply_markup=main_menu()
        )


# ============================================================
# HTTP SERVER
# ============================================================

async def health(request):
    return web.Response(
        text="Zolog AI is running"
    )


async def root(request):
    return web.Response(
        text="Zolog AI"
    )


async def start_web_server():
    app = web.Application()

    app.router.add_get(
        "/",
        root
    )

    app.router.add_get(
        "/health",
        health
    )

    runner = web.AppRunner(app)

    await runner.setup()

    site = web.TCPSite(
        runner,
        "0.0.0.0",
        PORT
    )

    await site.start()

    logger.info(
        "HTTP server started on port %s",
        PORT
    )


# ============================================================
# MAIN
# ============================================================

async def main():
    init_db()

    if not BOT_TOKEN:
        raise RuntimeError(
            "BOT_TOKEN не задан."
        )

    logger.info(
        "Starting Zolog AI..."
    )

    await bot.delete_webhook(
        drop_pending_updates=True
    )

    await start_web_server()

    logger.info(
        "Bot polling started."
    )

    await dp.start_polling(
        bot
    )


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        pass
